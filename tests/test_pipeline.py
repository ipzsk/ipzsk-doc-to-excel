import sys
import argparse
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('pipeline', ROOT / 'scripts/doc_to_excel.py')
pipeline = importlib.util.module_from_spec(spec)
spec.loader.exec_module(pipeline)


def args(**values):
    return argparse.Namespace(**values)


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.work = self.root / 'work'

    def tearDown(self):
        self.temp.cleanup()

    def make_docx(self):
        from docx import Document
        doc = Document()
        doc.add_heading('背景', 1)
        doc.add_paragraph('客户需求：支持离线使用。')
        doc.add_paragraph('海边的石头呈浅灰色。')
        doc.add_paragraph('海边的石头呈浅灰色。')
        doc.add_paragraph('需求和计划同时出现。')
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '编号'
        table.cell(0, 1).text = '原文'
        table.cell(1, 0).text = '00123'
        table.cell(1, 1).text = '=HYPERLINK("https://example.invalid","原文")'
        path = self.root / '示例.docx'
        doc.save(path)
        return path

    def prepare(self, path):
        return pipeline.prepare(args(inputs=[str(path)], work=str(self.work), schema=None))

    def test_real_docx_roundtrip_and_formula_safety(self):
        from openpyxl import load_workbook
        path = self.make_docx()
        first = self.prepare(path)
        self.assertEqual(first['records'], 7)
        self.assertTrue(self.prepare(path)['cached'])
        state = pipeline.load(self.work)
        pending = list(pipeline.pending(state))
        # Duplicate paragraphs share classification work, not provenance rows.
        self.assertEqual(len([r for r in pending if '海边' in r['text']]), 1)
        payload = pipeline.batch(args(work=str(self.work), limit=12, excerpt_chars=360, max_chars=8000))
        self.assertLessEqual(len(pipeline.json_text(payload)), 8000)
        decisions = [{'id': r['id'], 'category': '其他', 'confidence': 0.9, 'reason': '合成测试分类'} for r in payload['records']]
        file = self.root / 'decisions.json'
        file.write_text(json.dumps(decisions, ensure_ascii=False))
        pipeline.apply(args(work=str(self.work), decisions=str(file)))
        self.assertEqual(pipeline.status(pipeline.load(self.work))['unique_pending'], 0)
        output = self.root / 'result.xlsx'
        result = pipeline.export(args(work=str(self.work), output=str(output), overwrite=False))
        self.assertEqual(result['review_rows'], 0)
        book = load_workbook(output)
        self.assertEqual(book.sheetnames, ['总览', '分类明细', '待确认', '原始表格'])
        self.assertEqual(book['分类明细'].max_row, 8)
        raw = book['原始表格']
        self.assertEqual(raw['D3'].value, '00123')
        self.assertTrue(raw['E3'].value.startswith('=HYPERLINK'))
        self.assertEqual(raw['E3'].data_type, 's')
        self.assertEqual(book['分类明细'].freeze_panes, 'A2')
        book.close()
        with self.assertRaises(ValueError):
            pipeline.export(args(work=str(self.work), output=str(output), overwrite=False))

    def test_invalid_apply_is_atomic_and_stale_cache_rejected(self):
        path = self.make_docx()
        self.prepare(path)
        state_path = self.work / 'state.json'
        before = state_path.read_bytes()
        rid = next(pipeline.pending(pipeline.load(self.work)))['id']
        decisions = self.root / 'invalid.json'
        decisions.write_text(json.dumps([{'id': rid, 'category': '其他', 'confidence': .9}, {'id': 'unknown', 'category': '其他', 'confidence': .9}]))
        with self.assertRaises(ValueError):
            pipeline.apply(args(work=str(self.work), decisions=str(decisions)))
        self.assertEqual(before, state_path.read_bytes())
        from docx import Document
        doc = Document(path)
        doc.add_paragraph('修改后的源文件')
        doc.save(path)
        with self.assertRaises(ValueError):
            self.prepare(path)

    def test_text_pdf_and_empty_page_are_both_preserved(self):
        from pypdf import PdfReader, PdfWriter
        # Minimal valid PDF content stream avoids a test-only PDF generator dependency.
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 40 700 Td (Action: prepare next plan.) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
        writer.add_blank_page(width=612, height=792)
        path = self.root / 'pages.pdf'
        writer.write(path)
        self.assertEqual(len(PdfReader(path).pages), 2)
        self.prepare(path)
        state = pipeline.load(self.work)
        self.assertEqual(len(state['records']), 2)
        self.assertEqual(state['records'][1]['kind'], '待提取')
        output = self.root / 'pdf.xlsx'
        result = pipeline.export(args(work=str(self.work), output=str(output), overwrite=False))
        self.assertGreaterEqual(result['review_rows'], 1)

    def test_long_content_has_no_loss_and_batch_is_bounded(self):
        text = '甲乙丙丁' * 10000
        rows = list(pipeline.split_record({'text': text, 'location': '第 1 页', 'section': '', 'kind': 'PDF 页面'}))
        self.assertEqual(''.join(r['text'] for r in rows), text)
        self.assertTrue(all(len(r['text']) <= 12000 for r in rows))
        self.prepare(self.make_docx())
        batch = pipeline.batch(args(work=str(self.work), limit=1, excerpt_chars=100, max_chars=1000))
        self.assertEqual(len(batch['records']), 1)
        self.assertLessEqual(len(pipeline.json_text(batch)), 1000)
        with self.assertRaises(ValueError):
            pipeline.validate_schema({'categories': [{'name': 'x'}, {'name': 'x'}]})

    def test_budget_exhaustion_keeps_unresolved_records(self):
        self.prepare(self.make_docx())
        state = pipeline.load(self.work)
        state['shown_chars'] = 23950
        pipeline.atomic_json(self.work / 'state.json', state)
        result = pipeline.batch(args(work=str(self.work), limit=12, excerpt_chars=360, max_chars=8000))
        self.assertTrue(result['budget_exhausted'])
        self.assertGreater(result['unique_pending'], 0)
        self.assertEqual(pipeline.load(self.work)['shown_chars'], 23950)

    def test_low_confidence_is_reviewed(self):
        self.prepare(self.make_docx())
        state = pipeline.load(self.work)
        decisions = [{'id': r['id'], 'category': '其他', 'confidence': .3, 'reason': '信息不足'} for r in pipeline.pending(state)]
        file = self.root / 'low.json'
        file.write_text(json.dumps(decisions))
        pipeline.apply(args(work=str(self.work), decisions=str(file)))
        result = pipeline.export(args(work=str(self.work), output=str(self.root / 'review.xlsx'), overwrite=False))
        self.assertGreater(result['review_rows'], 0)


if __name__ == '__main__':
    unittest.main()

class OcrTests(unittest.TestCase):
    @unittest.skipUnless(sys.platform == 'darwin', 'macOS Vision OCR')
    def test_image_only_word_and_scanned_pdf(self):
        from PIL import Image, ImageDraw, ImageFont
        from docx import Document
        from reportlab.pdfgen import canvas
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            picture = root / 'scan.png'
            image = Image.new('RGB', (1400, 500), 'white')
            draw = ImageDraw.Draw(image)
            font = ImageFont.truetype('/System/Library/Fonts/Helvetica.ttc', 58)
            draw.text((50, 80), 'Invoice total 1234', font=font, fill='black')
            image.save(picture)
            word = root / 'image-only.docx'
            doc = Document(); doc.add_picture(str(picture)); doc.save(word)
            rows = list(pipeline.extract_docx(word))
            self.assertEqual(len(rows), 1)
            self.assertIn('1234', rows[0]['text'])
            self.assertTrue(rows[0]['needs_review'])
            pdf = root / 'scan.pdf'
            c = canvas.Canvas(str(pdf)); c.drawImage(str(picture), 30, 430, width=540, height=193); c.save()
            pages = list(pipeline.extract_pdf(pdf))
            self.assertIn('1234', pages[0]['text'])
            self.assertEqual(pages[0]['kind'], 'PDF OCR')
